import hashlib
import re
import zipfile
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document

MAX_BYTES = 10 * 1024 * 1024


def extract_resume(data: bytes, suffix: str) -> str:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Resume must be between 1 byte and 10 MB")
    try:
        if suffix.lower() == ".pdf":
            if not data.startswith(b"%PDF"):
                raise ValueError("Invalid PDF")
            with pymupdf.open(stream=data, filetype="pdf") as doc:
                if doc.needs_pass or len(doc) > 50:
                    raise ValueError("Encrypted or over-50-page PDF")
                text = "\n".join(page.get_text() for page in doc)
        elif suffix.lower() == ".docx":
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 30_000_000:
                    raise ValueError("DOCX expanded size too large")
            docx = Document(BytesIO(data))
            text = "\n".join(
                [p.text for p in docx.paragraphs]
                + [cell.text for table in docx.tables for row in table.rows for cell in row.cells]
            )
        else:
            raise ValueError("Only PDF and DOCX resumes are supported")
    except Exception as exc:
        raise ValueError("Could not read resume; use an unencrypted PDF or DOCX") from exc
    if len(text.strip()) < 10:
        raise ValueError("No usable text; OCR scanned resumes before uploading")
    return text[:200000]


def store_resume(data: bytes, suffix: str, folder: Path) -> tuple[Path, str]:
    content_hash = hashlib.sha256(data).hexdigest()
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{content_hash}{suffix.lower()}"
    path.write_bytes(data)
    return path, content_hash


def suggest_profile(text: str) -> dict:
    known = [
        "Python",
        "SQL",
        "FastAPI",
        "PyTorch",
        "TensorFlow",
        "Azure",
        "AWS",
        "Docker",
        "Kubernetes",
        "LLM",
        "RAG",
        "LangChain",
        "LiveKit",
        "Machine Learning",
    ]
    return {
        "name": "",
        "email": "",
        "skills": [
            s for s in known if re.search(r"(?<!\w)" + re.escape(s) + r"(?!\w)", text, re.I)
        ],
        "facts": [],
        "verified": False,
    }
